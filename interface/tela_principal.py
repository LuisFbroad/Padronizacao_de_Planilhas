import customtkinter as ctk
import tkinter as tk

from tkinter import filedialog, messagebox

from sqlalchemy import func, extract

from openpyxl import Workbook

from openpyxl.styles import Font, PatternFill, Alignment

from interface.tema import CORES

# Compatibilidade com versões do tema/componentes que acrescentaram estas cores.
CORES.setdefault("texto_botao", CORES.get("sidebar", "#071126"))
CORES.setdefault("destaque_claro", CORES.get("destaque", "#B08D3C"))
CORES.setdefault("texto_desativado", CORES.get("texto_secundario", "#B8C2D6"))

from interface.componentes import CardIndicador, Mensagem

from src.database.repository import (
    listar_anos_disponiveis,
    listar_areas,
    listar_ranking,
    contar_total_pessoas,
)

from src.database.connection import SessionLocal

from src.database.models import Area, Pessoa, Indicacao, Processo


class TelaPrincipal(ctk.CTkFrame):

    def __init__(self, master, **kwargs):

        super().__init__(master, fg_color=CORES["fundo"], **kwargs)

        self.areas = []

        self.area_selecionada = None

        self.tipo_indicacao_selecionado = "Todas"

        self.tela_atual = None

        self.criar_interface()

        self.carregar_areas()

        self.mostrar_dashboard()

    # =========================================================

    # ESTRUTURA PRINCIPAL

    # =========================================================

    def criar_interface(self):
        """Monta a estrutura executiva: navegação lateral e área de conteúdo."""
        self.configure(fg_color=CORES["fundo"])
        self.botoes_menu = {}

        self.sidebar = ctk.CTkFrame(
            self, width=238, fg_color=CORES["sidebar"], corner_radius=0
        )
        self.sidebar.pack(side="left", fill="y")
        self.sidebar.pack_propagate(False)

        marca = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        marca.pack(fill="x", padx=24, pady=(30, 34))
        ctk.CTkFrame(
            marca, width=5, height=46, corner_radius=3, fg_color=CORES["destaque"]
        ).pack(side="left", padx=(0, 12))
        textos_marca = ctk.CTkFrame(marca, fg_color="transparent")
        textos_marca.pack(side="left", fill="x", expand=True)
        ctk.CTkLabel(
            textos_marca,
            text="GAGC",
            text_color=CORES["texto"],
            font=ctk.CTkFont(size=27, weight="bold"),
            anchor="w",
        ).pack(anchor="w")
        ctk.CTkLabel(
            textos_marca,
            text="GESTÃO DE INDICAÇÕES",
            text_color=CORES["destaque"],
            font=ctk.CTkFont(size=9, weight="bold"),
            anchor="w",
        ).pack(anchor="w", pady=(2, 0))

        ctk.CTkLabel(
            self.sidebar,
            text="NAVEGAÇÃO",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=10, weight="bold"),
            anchor="w",
        ).pack(fill="x", padx=24, pady=(0, 10))
        self.criar_botao_menu(
            "▦   Dashboard", self.mostrar_dashboard, chave="Dashboard"
        )
        self.criar_botao_menu(
            "▥   Indicadores", self.mostrar_indicadores, chave="Indicadores"
        )
        self.criar_botao_menu(
            "▤   Relatórios", self.mostrar_relatorios, chave="Relatórios"
        )

        rodape = ctk.CTkFrame(self.sidebar, fg_color="transparent")
        rodape.pack(side="bottom", fill="x", padx=22, pady=22)
        ctk.CTkFrame(rodape, height=1, fg_color=CORES["borda"]).pack(
            fill="x", pady=(0, 13)
        )
        ctk.CTkLabel(
            rodape,
            text="SISTEMA INTERNO",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=10, weight="bold"),
            anchor="w",
        ).pack(anchor="w")
        ctk.CTkLabel(
            rodape,
            text="GAGC  •  Área administrativa",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=10),
            anchor="w",
        ).pack(anchor="w", pady=(4, 0))

        self.area_principal = ctk.CTkFrame(
            self, fg_color=CORES["fundo"], corner_radius=0
        )
        self.area_principal.pack(side="left", fill="both", expand=True)

    def criar_botao_menu(self, texto, comando, chave=None):
        chave = chave or texto
        botao = ctk.CTkButton(
            self.sidebar,
            text=texto,
            command=comando,
            height=48,
            corner_radius=9,
            fg_color="transparent",
            hover_color=CORES["sidebar_hover"],
            text_color=CORES["texto_secundario"],
            anchor="w",
            font=ctk.CTkFont(size=13),
            border_spacing=14,
        )
        botao.pack(padx=13, pady=4, fill="x")
        self.botoes_menu[chave] = botao
        return botao

    def limpar_area(self):
        for widget in self.area_principal.winfo_children():
            widget.destroy()

    def criar_cabecalho(self, titulo, subtitulo=""):
        frame = ctk.CTkFrame(self.area_principal, fg_color="transparent")
        frame.pack(fill="x", padx=36, pady=(30, 20))
        linha = ctk.CTkFrame(frame, fg_color="transparent")
        linha.pack(fill="x")
        ctk.CTkLabel(
            linha,
            text=titulo,
            text_color=CORES["texto"],
            font=ctk.CTkFont(size=29, weight="bold"),
            anchor="w",
        ).pack(side="left", anchor="w")
        ctk.CTkLabel(
            linha,
            text="GAGC  /  SISTEMA DE INDICAÇÕES",
            text_color=CORES["destaque"],
            font=ctk.CTkFont(size=10, weight="bold"),
            anchor="e",
        ).pack(side="right", padx=(12, 0))
        if subtitulo:
            ctk.CTkLabel(
                frame,
                text=subtitulo,
                text_color=CORES["texto_secundario"],
                font=ctk.CTkFont(size=13),
                anchor="w",
            ).pack(anchor="w", pady=(6, 0))
        ctk.CTkFrame(frame, height=1, fg_color=CORES["borda"]).pack(
            fill="x", pady=(17, 0)
        )

    def definir_menu_ativo(self, tela):
        nomes = {
            "dashboard": "Dashboard",
            "indicadores": "Indicadores",
            "relatorios": "Relatórios",
        }
        ativo = nomes.get(tela)
        for nome, botao in getattr(self, "botoes_menu", {}).items():
            selecionado = nome == ativo
            botao.configure(
                fg_color=CORES["sidebar_hover"] if selecionado else "transparent",
                text_color=(
                    CORES["destaque"] if selecionado else CORES["texto_secundario"]
                ),
                font=ctk.CTkFont(size=13, weight="bold" if selecionado else "normal"),
            )

    def criar_painel(self, parent, titulo, subtitulo=None):
        painel = ctk.CTkFrame(
            parent,
            fg_color=CORES["card"],
            corner_radius=14,
            border_width=1,
            border_color=CORES["borda"],
        )
        cab = ctk.CTkFrame(painel, fg_color="transparent")
        cab.pack(fill="x", padx=20, pady=(18, 12))
        ctk.CTkLabel(
            cab,
            text=titulo,
            text_color=CORES["texto"],
            font=ctk.CTkFont(size=16, weight="bold"),
            anchor="w",
        ).pack(anchor="w")
        if subtitulo:
            ctk.CTkLabel(
                cab,
                text=subtitulo,
                text_color=CORES["texto_secundario"],
                font=ctk.CTkFont(size=11),
                anchor="w",
            ).pack(anchor="w", pady=(4, 0))
        return painel

    def criar_grafico_barras(
        self, parent, dados, titulo="Indicações por pessoa", limite=8
    ):
        """Cria gráfico de barras horizontal em Canvas, sem dependência externa."""
        painel = self.criar_painel(
            parent, titulo, "Comparativo dos nomes com maior quantidade"
        )
        canvas = tk.Canvas(
            painel,
            height=300,
            bg=CORES["card"],
            highlightthickness=0,
            bd=0,
            relief="flat",
        )
        canvas.pack(fill="both", expand=True, padx=14, pady=(0, 16))
        dados = list(dados or [])[:limite]
        if titulo == "Destaques do ranking":
            self.dados_grafico_dashboard = dados

        def desenhar(_evento=None):
            canvas.delete("all")
            dados_atuais = (
                getattr(self, "dados_grafico_dashboard", dados)
                if titulo == "Destaques do ranking"
                else dados
            )
            largura = max(canvas.winfo_width(), 320)
            altura = max(canvas.winfo_height(), 220)
            if not dados_atuais:
                canvas.create_text(
                    largura / 2,
                    altura / 2,
                    text="Sem dados para exibir",
                    fill=CORES["texto_secundario"],
                    font=("Segoe UI", 11),
                )
                return
            max_valor = max([int(item[1] or 0) for item in dados_atuais] + [1])
            margem_esquerda = min(150, max(95, int(largura * 0.34)))
            margem_direita = 42
            topo = 10
            altura_linha = max(27, min(38, (altura - 20) // max(len(dados_atuais), 1)))
            area_barra = max(40, largura - margem_esquerda - margem_direita)
            for i, (nome, valor) in enumerate(dados_atuais):
                y = topo + i * altura_linha
                rotulo = str(nome or "Sem nome")
                if len(rotulo) > 21:
                    rotulo = rotulo[:19] + "…"
                canvas.create_text(
                    margem_esquerda - 12,
                    y + 9,
                    text=rotulo,
                    anchor="e",
                    fill=CORES["texto_secundario"],
                    font=("Segoe UI", 9),
                )
                canvas.create_rectangle(
                    margem_esquerda,
                    y,
                    margem_esquerda + area_barra,
                    y + 17,
                    fill=CORES["fundo_secundario"],
                    outline="",
                )
                barra = area_barra * max(0, int(valor or 0)) / max_valor
                canvas.create_rectangle(
                    margem_esquerda,
                    y,
                    margem_esquerda + barra,
                    y + 17,
                    fill=CORES["destaque"],
                    outline="",
                )
                canvas.create_text(
                    min(largura - 8, margem_esquerda + area_barra + 8),
                    y + 8,
                    text=str(valor),
                    anchor="w",
                    fill=CORES["texto"],
                    font=("Segoe UI", 9, "bold"),
                )

        canvas.bind("<Configure>", desenhar)
        desenhar()
        painel.pack(fill="both", expand=True)
        if titulo == "Destaques do ranking":
            self.canvas_grafico_dashboard = canvas
        return painel, canvas

    def carregar_areas(self):

        try:

            self.areas = listar_areas()

            if self.areas:

                self.area_selecionada = self.areas[0]

        except Exception as erro:

            messagebox.showerror(
                "Erro",
                f"Não foi possível carregar as áreas:\n\n{erro}",
            )

    # =========================================================

    # DASHBOARD

    # =========================================================

    def mostrar_dashboard(self):
        self.tela_atual = "dashboard"
        self.definir_menu_ativo("dashboard")
        self.limpar_area()
        self.criar_cabecalho(
            "Dashboard", "Visão geral das indicações e desempenho por área."
        )

        conteudo = ctk.CTkFrame(self.area_principal, fg_color="transparent")
        conteudo.pack(fill="both", expand=True, padx=36, pady=(0, 28))

        cards = ctk.CTkFrame(conteudo, fg_color="transparent")
        cards.pack(fill="x", pady=(0, 18))
        cards.grid_columnconfigure((0, 1), weight=1, uniform="dashboard_cards")
        self.card_processos = CardIndicador(
            cards, titulo="Processos", valor=self.obter_total_processos(), icone="▣"
        )
        self.card_processos.grid(row=0, column=0, padx=(0, 9), sticky="ew")
        self.card_indicacoes_dashboard = CardIndicador(
            cards,
            titulo="Indicadores",
            valor=self.obter_total_indicacoes_dashboard(),
            icone="↗",
        )
        self.card_indicacoes_dashboard.grid(row=0, column=1, padx=(9, 0), sticky="ew")

        self.criar_filtros_dashboard(parent=conteudo)

        area_inferior = ctk.CTkFrame(conteudo, fg_color="transparent")
        area_inferior.pack(fill="both", expand=True, pady=(17, 0))
        area_inferior.grid_columnconfigure(0, weight=5, uniform="dash_bottom")
        area_inferior.grid_columnconfigure(1, weight=6, uniform="dash_bottom")
        area_inferior.grid_rowconfigure(0, weight=1)

        grafico_host = ctk.CTkFrame(area_inferior, fg_color="transparent")
        self.grafico_dashboard_host = grafico_host
        grafico_host.grid(row=0, column=0, sticky="nsew", padx=(0, 9))
        ranking_host = ctk.CTkFrame(area_inferior, fg_color="transparent")
        ranking_host.grid(row=0, column=1, sticky="nsew", padx=(9, 0))
        self.criar_grafico_barras(grafico_host, [], "Destaques do ranking")
        self.criar_area_ranking(parent=ranking_host)
        self.atualizar_ranking_banco()

    def criar_filtros_dashboard(self, parent=None):
        parent = parent or self.area_principal
        container = ctk.CTkFrame(
            parent,
            fg_color=CORES["card"],
            corner_radius=14,
            border_width=1,
            border_color=CORES["borda"],
        )
        container.pack(fill="x")
        ctk.CTkLabel(
            container,
            text="FILTROS DO DASHBOARD",
            text_color=CORES["destaque"],
            font=ctk.CTkFont(size=10, weight="bold"),
        ).pack(anchor="w", padx=20, pady=(15, 3))
        filtros = ctk.CTkFrame(container, fg_color="transparent")
        filtros.pack(fill="x", padx=20, pady=(6, 18))
        filtros.grid_columnconfigure((0, 1), weight=1, uniform="dashboard_filters")

        ctk.CTkLabel(
            filtros,
            text="Área de análise",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=11, weight="bold"),
        ).grid(row=0, column=0, sticky="w", padx=(0, 10), pady=(0, 6))
        self.combo_areas = ctk.CTkComboBox(
            filtros,
            values=[area.nome for area in self.areas],
            height=40,
            command=self.selecionar_area,
            corner_radius=8,
            fg_color=CORES["fundo_secundario"],
            border_color=CORES["borda"],
            button_color=CORES["destaque"],
            button_hover_color=CORES["destaque_hover"],
        )
        self.combo_areas.grid(row=1, column=0, sticky="ew", padx=(0, 10))
        if self.area_selecionada:
            self.combo_areas.set(self.area_selecionada.nome)
        elif self.areas:
            self.combo_areas.set(self.areas[0].nome)

        ctk.CTkLabel(
            filtros,
            text="Tipo de indicação",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=11, weight="bold"),
        ).grid(row=0, column=1, sticky="w", padx=(10, 0), pady=(0, 6))
        self.tipos_indicacao = [
            "Todas",
            "Primária",
            "Secundária",
            "Terciária",
            "Quaternária",
            "Final",
        ]
        self.combo_tipos = ctk.CTkComboBox(
            filtros,
            values=self.tipos_indicacao,
            height=40,
            command=self.selecionar_tipo_indicacao,
            corner_radius=8,
            fg_color=CORES["fundo_secundario"],
            border_color=CORES["borda"],
            button_color=CORES["destaque"],
            button_hover_color=CORES["destaque_hover"],
        )
        self.combo_tipos.grid(row=1, column=1, sticky="ew", padx=(10, 0))
        self.combo_tipos.set(self.tipo_indicacao_selecionado)

    def criar_area_ranking(self, parent=None):
        parent = parent or self.area_principal
        painel = self.criar_painel(
            parent, "Ranking de indicações", "Pessoas com maior volume de indicações"
        )
        painel.pack(fill="both", expand=True)
        self.ranking_scroll = ctk.CTkScrollableFrame(
            painel, fg_color="transparent", corner_radius=0
        )
        self.ranking_scroll.pack(fill="both", expand=True, padx=14, pady=(0, 15))

    def selecionar_area(self, nome_area):

        for area in self.areas:

            if area.nome == nome_area:

                self.area_selecionada = area

                break

        self.atualizar_cards()

        self.atualizar_ranking_banco()

    def selecionar_tipo_indicacao(self, tipo):

        self.tipo_indicacao_selecionado = tipo

        self.atualizar_cards()

        self.atualizar_ranking_banco()

    def atualizar_cards(self):

        if not hasattr(self, "card_processos"):

            return

        self.card_processos.atualizar_valor(self.obter_total_processos())
        self.card_indicacoes_dashboard.atualizar_valor(
            self.obter_total_indicacoes_dashboard()
        )

    def obter_total_indicacoes_dashboard(self):
        """Conta indicações da área e do tipo selecionados no Dashboard."""
        if self.area_selecionada is None:
            return "0"
        if "projet" in (self.area_selecionada.nome or "").strip().casefold():
            return "0"
        session = SessionLocal()
        try:
            consulta = session.query(func.count(Indicacao.id)).filter(
                Indicacao.area_id == self.area_selecionada.id
            )
            if self.tipo_indicacao_selecionado != "Todas":
                consulta = consulta.filter(
                    Indicacao.tipo_indicacao == self.tipo_indicacao_selecionado
                )
            return str(consulta.scalar() or 0)
        except Exception as erro:
            print(f"Erro ao contar indicações do Dashboard: {erro}")
            return "0"
        finally:
            session.close()

    def obter_total_processos(self):

        session = SessionLocal()

        try:

            if self.area_selecionada is None:

                return "0"

            nome_area = (self.area_selecionada.nome or "").strip().casefold()

            if "projet" in nome_area:

                consulta = session.query(
                    func.coalesce(func.max(Processo.quantidade), 0)
                ).filter(Processo.area_id == self.area_selecionada.id)

                return str(consulta.scalar() or 0)

            consulta = session.query(
                func.coalesce(func.max(Indicacao.quantidade), 0)
            ).filter(Indicacao.area_id == self.area_selecionada.id)

            if self.tipo_indicacao_selecionado != "Todas":

                consulta = consulta.filter(
                    Indicacao.tipo_indicacao == self.tipo_indicacao_selecionado
                )

            return str(consulta.scalar() or 0)

        except Exception as erro:

            print(f"Erro ao contar processos: {erro}")

            return "0"

        finally:

            session.close()

    def obter_total_pessoas(self):

        try:

            return str(
                contar_total_pessoas(
                    area_id=(
                        self.area_selecionada.id if self.area_selecionada else None
                    ),
                    tipo_indicacao=self.tipo_indicacao_selecionado,
                )
            )

        except Exception:

            return "0"

    def atualizar_ranking_banco(self):
        if not hasattr(self, "ranking_scroll"):
            return
        for widget in self.ranking_scroll.winfo_children():
            widget.destroy()
        if self.area_selecionada is None:
            Mensagem(
                self.ranking_scroll, "Nenhuma área selecionada.", tipo="alerta"
            ).pack(pady=30)
            return
        try:
            ranking = listar_ranking(
                self.area_selecionada.id, self.tipo_indicacao_selecionado
            )
        except Exception as erro:
            Mensagem(
                self.ranking_scroll, f"Erro ao consultar o banco:\n{erro}", tipo="erro"
            ).pack(pady=30)
            return
        self.dados_grafico_dashboard = list(ranking or [])[:8]
        canvas_grafico = getattr(self, "canvas_grafico_dashboard", None)
        if canvas_grafico and canvas_grafico.winfo_exists():
            canvas_grafico.event_generate("<Configure>")
        if not ranking:
            Mensagem(
                self.ranking_scroll,
                "Nenhuma indicação encontrada para os filtros selecionados.",
                tipo="alerta",
            ).pack(pady=30)
            return

        for indice, (nome, quantidade) in enumerate(ranking):
            linha = ctk.CTkFrame(
                self.ranking_scroll,
                fg_color=(
                    CORES["fundo_secundario"] if indice % 2 == 0 else CORES["card"]
                ),
                corner_radius=9,
                border_width=1,
                border_color=CORES["borda"],
            )
            linha.pack(fill="x", pady=4, padx=2)
            ctk.CTkLabel(
                linha,
                text=f"{indice + 1:02d}",
                width=42,
                text_color=CORES["destaque"],
                font=ctk.CTkFont(size=12, weight="bold"),
            ).pack(side="left", padx=(10, 4), pady=11)
            ctk.CTkLabel(
                linha,
                text=str(nome),
                text_color=CORES["texto"],
                font=ctk.CTkFont(size=12, weight="bold"),
                anchor="w",
            ).pack(side="left", fill="x", expand=True, padx=8, pady=11)
            ctk.CTkLabel(
                linha,
                text=f"{quantidade}",
                text_color=CORES["destaque"],
                font=ctk.CTkFont(size=12, weight="bold"),
            ).pack(side="right", padx=(8, 14), pady=11)

    def mostrar_indicadores(self):
        self.tela_atual = "indicadores"
        self.definir_menu_ativo("indicadores")
        self.limpar_area()
        self.criar_cabecalho(
            "Indicadores",
            "Acompanhe o desempenho consolidado de todas as áreas por período.",
        )

        container = ctk.CTkFrame(self.area_principal, fg_color="transparent")
        container.pack(fill="both", expand=True, padx=36, pady=(0, 28))

        filtros = ctk.CTkFrame(
            container,
            fg_color=CORES["card"],
            corner_radius=14,
            border_width=1,
            border_color=CORES["borda"],
        )
        filtros.pack(fill="x", pady=(0, 18))
        for coluna in range(4):
            filtros.grid_columnconfigure(
                coluna, weight=1, uniform="indicadores_filters"
            )
        ctk.CTkLabel(
            filtros,
            text="FILTROS DE PERÍODO",
            text_color=CORES["destaque"],
            font=ctk.CTkFont(size=10, weight="bold"),
        ).grid(row=0, column=0, columnspan=4, sticky="w", padx=20, pady=(15, 8))

        for coluna, titulo in enumerate(["Ano", "Mês", "Semana", "Tipo de indicação"]):
            ctk.CTkLabel(
                filtros,
                text=titulo,
                text_color=CORES["texto_secundario"],
                font=ctk.CTkFont(size=11, weight="bold"),
            ).grid(row=1, column=coluna, padx=10, pady=(0, 6), sticky="w")

        self.meses_indicadores = {
            "Todos": None,
            "Janeiro": 1,
            "Fevereiro": 2,
            "Março": 3,
            "Abril": 4,
            "Maio": 5,
            "Junho": 6,
            "Julho": 7,
            "Agosto": 8,
            "Setembro": 9,
            "Outubro": 10,
            "Novembro": 11,
            "Dezembro": 12,
        }
        anos = listar_anos_disponiveis()
        self.combo_ano_indicadores = ctk.CTkComboBox(
            filtros,
            values=["Todos"] + [str(ano) for ano in anos],
            command=self.alterar_filtro_indicadores,
            height=39,
            corner_radius=8,
        )
        self.combo_ano_indicadores.grid(
            row=2, column=0, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_ano_indicadores.set("Todos")
        self.combo_mes_indicadores = ctk.CTkComboBox(
            filtros,
            values=list(self.meses_indicadores.keys()),
            command=self.alterar_filtro_indicadores,
            height=39,
            corner_radius=8,
        )
        self.combo_mes_indicadores.grid(
            row=2, column=1, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_mes_indicadores.set("Todos")
        self.combo_semana_indicadores = ctk.CTkComboBox(
            filtros,
            values=["Todas", "1ª", "2ª", "3ª", "4ª"],
            command=self.alterar_filtro_indicadores,
            height=39,
            corner_radius=8,
        )
        self.combo_semana_indicadores.grid(
            row=2, column=2, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_semana_indicadores.set("Todas")
        self.combo_tipo_indicadores = ctk.CTkComboBox(
            filtros,
            values=[
                "Todas",
                "Primária",
                "Secundária",
                "Terciária",
                "Quaternária",
                "Final",
            ],
            command=self.alterar_tipo_indicadores,
            height=39,
            corner_radius=8,
        )
        self.combo_tipo_indicadores.grid(
            row=2, column=3, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_tipo_indicadores.set(
            self.tipo_indicacao_selecionado
            if self.tipo_indicacao_selecionado != "Todas"
            else "Todas"
        )

        cards = ctk.CTkFrame(container, fg_color="transparent")
        cards.pack(fill="x", pady=(0, 18))
        cards.grid_columnconfigure((0, 1), weight=1, uniform="indicator_cards")
        self.card_indicadores_processos = CardIndicador(
            cards, titulo="Processos no período", valor="0", icone="▣"
        )
        self.card_indicadores_processos.grid(row=0, column=0, padx=(0, 9), sticky="ew")
        self.card_indicadores_pessoas = CardIndicador(
            cards, titulo="Pessoas com indicações", valor="0", icone="↗"
        )
        self.card_indicadores_pessoas.grid(row=0, column=1, padx=(9, 0), sticky="ew")

        area_inferior = ctk.CTkFrame(container, fg_color="transparent")
        area_inferior.pack(fill="both", expand=True)
        area_inferior.grid_columnconfigure(0, weight=5, uniform="indicator_bottom")
        area_inferior.grid_columnconfigure(1, weight=6, uniform="indicator_bottom")
        area_inferior.grid_rowconfigure(0, weight=1)
        self.grafico_indicadores_host = ctk.CTkFrame(
            area_inferior, fg_color="transparent"
        )
        self.grafico_indicadores_host.grid(row=0, column=0, sticky="nsew", padx=(0, 9))
        ranking_host = ctk.CTkFrame(area_inferior, fg_color="transparent")
        ranking_host.grid(row=0, column=1, sticky="nsew", padx=(9, 0))
        painel = self.criar_painel(
            ranking_host,
            "Ranking geral",
            "Todas as áreas • filtrado pelo período selecionado",
        )
        painel.pack(fill="both", expand=True)
        self.ranking_indicadores = ctk.CTkScrollableFrame(
            painel, fg_color="transparent", corner_radius=0
        )
        self.ranking_indicadores.pack(fill="both", expand=True, padx=14, pady=(0, 15))
        self.atualizar_indicadores()

    def obter_periodo_indicadores(self):

        ano_texto = self.combo_ano_indicadores.get()

        mes_texto = self.combo_mes_indicadores.get()

        semana_texto = self.combo_semana_indicadores.get()

        ano = int(ano_texto) if ano_texto != "Todos" else None

        mes = self.meses_indicadores.get(mes_texto)

        semanas = {
            "1ª": 1,
            "2ª": 2,
            "3ª": 3,
            "4ª": 4,
        }

        semana = semanas.get(semana_texto)

        return ano, mes, semana

    def alterar_filtro_indicadores(self, valor=None):
        """Atualiza os indicadores quando ano, mês ou semana forem alterados."""
        self.atualizar_indicadores()

    def alterar_tipo_indicadores(self, valor=None):
        """Atualiza o tipo de indicação usado na aba Indicadores.

        Este filtro é independente da área selecionada no Dashboard.
        """
        if valor is None and hasattr(self, "combo_tipo_indicadores"):
            valor = self.combo_tipo_indicadores.get()

        self.tipo_indicacao_selecionado = valor or "Todas"
        self.atualizar_indicadores()

    def aplicar_filtros_periodo(
        self,
        consulta,
        ano=None,
        mes=None,
        semana=None,
        campo_data=None,
    ):

        data = campo_data or Indicacao.data_indicacao

        if ano is not None:

            consulta = consulta.filter(extract("year", data) == ano)

        if mes is not None:

            consulta = consulta.filter(extract("month", data) == mes)

        if semana is not None:

            limites = {
                1: (1, 7),
                2: (8, 14),
                3: (15, 21),
                4: (22, 31),
            }

            inicio, fim = limites[semana]

            dia = extract("day", data)

            consulta = consulta.filter(
                dia >= inicio,
                dia <= fim,
            )

        return consulta

    def aplicar_filtros_indicadores(self, consulta):
        """Aplica filtros da aba Indicadores sem herdar a área do Dashboard."""
        if self.tipo_indicacao_selecionado != "Todas":
            consulta = consulta.filter(
                Indicacao.tipo_indicacao == self.tipo_indicacao_selecionado
            )

        return consulta

    def contar_processos_periodo(self, ano=None, mes=None, semana=None):
        """Soma os processos de todas as áreas no período selecionado."""
        session = SessionLocal()

        try:
            total = 0
            areas = session.query(Area).order_by(Area.id).all()

            for area in areas:
                nome_area = (area.nome or "").strip().casefold()

                if "projet" in nome_area:
                    # O tipo de indicação não se aplica à área de projetos.
                    if self.tipo_indicacao_selecionado != "Todas":
                        continue

                    consulta = session.query(
                        func.coalesce(func.max(Processo.quantidade), 0)
                    ).filter(Processo.area_id == area.id)

                    consulta = self.aplicar_filtros_periodo(
                        consulta,
                        ano=ano,
                        mes=mes,
                        semana=semana,
                        campo_data=Processo.data_processo,
                    )
                else:
                    consulta = session.query(
                        func.coalesce(func.max(Indicacao.quantidade), 0)
                    ).filter(Indicacao.area_id == area.id)

                    consulta = self.aplicar_filtros_indicadores(consulta)
                    consulta = self.aplicar_filtros_periodo(
                        consulta, ano=ano, mes=mes, semana=semana
                    )

                total += int(consulta.scalar() or 0)

            return total
        except Exception as erro:
            print(f"Erro ao contar processos de todas as áreas: {erro}")
            return 0
        finally:
            session.close()

    def contar_pessoas_periodo(self, ano=None, mes=None, semana=None):

        session = SessionLocal()

        try:

            consulta = session.query(func.count(func.distinct(Indicacao.pessoa_id)))

            consulta = self.aplicar_filtros_indicadores(consulta)

            consulta = self.aplicar_filtros_periodo(consulta, ano, mes, semana)

            return consulta.scalar() or 0

        finally:

            session.close()

    def obter_ranking_indicadores(
        self,
        ano=None,
        mes=None,
        semana=None,
    ):

        session = SessionLocal()

        try:

            consulta = session.query(
                Pessoa.nome,
                func.count(Indicacao.id).label("total_indicacoes"),
            ).join(
                Indicacao,
                Indicacao.pessoa_id == Pessoa.id,
            )

            consulta = self.aplicar_filtros_indicadores(consulta)

            consulta = self.aplicar_filtros_periodo(consulta, ano, mes, semana)

            return (
                consulta.group_by(Pessoa.id, Pessoa.nome)
                .order_by(
                    func.count(Indicacao.id).desc(),
                    Pessoa.nome,
                )
                .all()
            )

        finally:

            session.close()

    def atualizar_indicadores(self):
        if not hasattr(self, "ranking_indicadores"):
            return
        for widget in self.ranking_indicadores.winfo_children():
            widget.destroy()
        ano, mes, semana = self.obter_periodo_indicadores()
        try:
            total_processos = self.contar_processos_periodo(
                ano=ano, mes=mes, semana=semana
            )
            total_pessoas = self.contar_pessoas_periodo(ano=ano, mes=mes, semana=semana)
            self.card_indicadores_processos.atualizar_valor(total_processos)
            self.card_indicadores_pessoas.atualizar_valor(total_pessoas)
            ranking = self.obter_ranking_indicadores(ano=ano, mes=mes, semana=semana)

            # Recria o gráfico com o ranking atual, sempre consolidado entre as áreas.
            for widget in self.grafico_indicadores_host.winfo_children():
                widget.destroy()
            self.criar_grafico_barras(
                self.grafico_indicadores_host, ranking, "Indicações por pessoa"
            )

            if not ranking:
                Mensagem(
                    self.ranking_indicadores,
                    "Nenhuma indicação encontrada para esse período.",
                    tipo="alerta",
                ).pack(pady=30)
                return
            for posicao, (nome, quantidade) in enumerate(ranking):
                linha = ctk.CTkFrame(
                    self.ranking_indicadores,
                    fg_color=(
                        CORES["fundo_secundario"] if posicao % 2 == 0 else CORES["card"]
                    ),
                    corner_radius=9,
                    border_width=1,
                    border_color=CORES["borda"],
                )
                linha.pack(fill="x", pady=4, padx=2)
                ctk.CTkLabel(
                    linha,
                    text=f"{posicao + 1:02d}",
                    width=42,
                    text_color=CORES["destaque"],
                    font=ctk.CTkFont(size=12, weight="bold"),
                ).pack(side="left", padx=(10, 4), pady=11)
                ctk.CTkLabel(
                    linha,
                    text=str(nome),
                    text_color=CORES["texto"],
                    font=ctk.CTkFont(size=12, weight="bold"),
                    anchor="w",
                ).pack(side="left", fill="x", expand=True, padx=8, pady=11)
                ctk.CTkLabel(
                    linha,
                    text=f"{quantidade} indicações",
                    text_color=CORES["texto_secundario"],
                    font=ctk.CTkFont(size=11),
                ).pack(side="right", padx=(8, 14), pady=11)
        except Exception as erro:
            Mensagem(
                self.ranking_indicadores,
                f"Erro ao carregar indicadores:\n{erro}",
                tipo="erro",
            ).pack(pady=30)

    def mostrar_relatorios(self):
        self.tela_atual = "relatorios"
        self.definir_menu_ativo("relatorios")
        self.limpar_area()
        self.criar_cabecalho(
            "Relatórios por área",
            "Consulte os resultados, refine o período e exporte os dados para Excel.",
        )

        container = ctk.CTkFrame(self.area_principal, fg_color="transparent")
        container.pack(fill="both", expand=True, padx=36, pady=(0, 28))

        filtros = ctk.CTkFrame(
            container,
            fg_color=CORES["card"],
            corner_radius=14,
            border_width=1,
            border_color=CORES["borda"],
        )
        filtros.pack(fill="x", pady=(0, 18))
        for coluna in range(4):
            filtros.grid_columnconfigure(coluna, weight=1, uniform="report_filters")
        ctk.CTkLabel(
            filtros,
            text="PARÂMETROS DO RELATÓRIO",
            text_color=CORES["destaque"],
            font=ctk.CTkFont(size=10, weight="bold"),
        ).grid(row=0, column=0, columnspan=4, sticky="w", padx=20, pady=(15, 8))
        for coluna, titulo in enumerate(["Ano", "Mês", "Semana", "Tipo de indicação"]):
            ctk.CTkLabel(
                filtros,
                text=titulo,
                text_color=CORES["texto_secundario"],
                font=ctk.CTkFont(size=11, weight="bold"),
            ).grid(row=1, column=coluna, padx=10, pady=(0, 6), sticky="w")

        self.meses_relatorios = {
            "Todos": None,
            "Janeiro": 1,
            "Fevereiro": 2,
            "Março": 3,
            "Abril": 4,
            "Maio": 5,
            "Junho": 6,
            "Julho": 7,
            "Agosto": 8,
            "Setembro": 9,
            "Outubro": 10,
            "Novembro": 11,
            "Dezembro": 12,
        }
        anos = listar_anos_disponiveis()
        estilo_combo = {
            "height": 39,
            "corner_radius": 8,
            "fg_color": CORES["fundo_secundario"],
            "border_color": CORES["borda"],
            "button_color": CORES["destaque"],
            "button_hover_color": CORES["destaque_hover"],
        }
        self.combo_ano_relatorios = ctk.CTkComboBox(
            filtros,
            values=["Todos"] + [str(ano) for ano in anos],
            command=self.alterar_filtro_relatorio,
            **estilo_combo,
        )
        self.combo_ano_relatorios.grid(
            row=2, column=0, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_ano_relatorios.set("Todos")
        self.combo_mes_relatorios = ctk.CTkComboBox(
            filtros,
            values=list(self.meses_relatorios.keys()),
            command=self.alterar_filtro_relatorio,
            **estilo_combo,
        )
        self.combo_mes_relatorios.grid(
            row=2, column=1, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_mes_relatorios.set("Todos")
        self.combo_semana_relatorios = ctk.CTkComboBox(
            filtros,
            values=["Todas", "1ª", "2ª", "3ª", "4ª"],
            command=self.alterar_filtro_relatorio,
            **estilo_combo,
        )
        self.combo_semana_relatorios.grid(
            row=2, column=2, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_semana_relatorios.set("Todas")
        self.combo_tipo_relatorios = ctk.CTkComboBox(
            filtros,
            values=[
                "Todas",
                "Primária",
                "Secundária",
                "Terciária",
                "Quaternária",
                "Final",
            ],
            command=self.alterar_filtro_relatorio,
            **estilo_combo,
        )
        self.combo_tipo_relatorios.grid(
            row=2, column=3, padx=10, pady=(0, 17), sticky="ew"
        )
        self.combo_tipo_relatorios.set("Todas")

        cards = ctk.CTkFrame(container, fg_color="transparent")
        cards.pack(fill="x", pady=(0, 18))
        cards.grid_columnconfigure((0, 1), weight=1, uniform="report_cards")
        self.card_relatorio_processos = CardIndicador(
            cards, titulo="Total de processos", valor="0", icone="▣"
        )
        self.card_relatorio_processos.grid(row=0, column=0, padx=(0, 9), sticky="ew")
        self.card_relatorio_areas = CardIndicador(
            cards, titulo="Áreas com resultados", valor="0", icone="▤"
        )
        self.card_relatorio_areas.grid(row=0, column=1, padx=(9, 0), sticky="ew")

        tabela_painel = self.criar_painel(
            container, "Resumo por área", "Resultados conforme os filtros selecionados"
        )
        tabela_painel.pack(fill="both", expand=True)
        barra_acoes = ctk.CTkFrame(tabela_painel, fg_color="transparent")
        barra_acoes.pack(fill="x", padx=18, pady=(0, 12))
        ctk.CTkLabel(
            barra_acoes,
            text="Os dados podem ser exportados em formato .xlsx",
            text_color=CORES["texto_secundario"],
            font=ctk.CTkFont(size=11),
        ).pack(side="left")
        ctk.CTkButton(
            barra_acoes,
            text="Exportar para Excel",
            command=self.exportar_relatorio_excel,
            height=38,
            corner_radius=8,
            fg_color=CORES["destaque"],
            hover_color=CORES["destaque_hover"],
            text_color=CORES.get("texto_botao", CORES["sidebar"]),
            font=ctk.CTkFont(size=12, weight="bold"),
        ).pack(side="right")

        self.tabela_relatorios = ctk.CTkFrame(tabela_painel, fg_color="transparent")
        self.tabela_relatorios.pack(fill="both", expand=True, padx=18, pady=(0, 18))
        self.tabela_relatorios.grid_columnconfigure(0, weight=3)
        self.tabela_relatorios.grid_columnconfigure(1, weight=1)
        self.tabela_relatorios.grid_columnconfigure(2, weight=1)
        self.tabela_relatorios.grid_rowconfigure(1, weight=1)

        cabecalho_tabela = ctk.CTkFrame(
            self.tabela_relatorios, fg_color=CORES["fundo_secundario"], corner_radius=7
        )
        cabecalho_tabela.grid(row=0, column=0, columnspan=3, sticky="ew", pady=(0, 6))
        for coluna, titulo in enumerate(["Área", "Processos", "Indicações"]):
            cabecalho_tabela.grid_columnconfigure(
                coluna, weight=3 if coluna == 0 else 1
            )
            ctk.CTkLabel(
                cabecalho_tabela,
                text=titulo,
                text_color=CORES["texto"],
                font=ctk.CTkFont(size=12, weight="bold"),
                anchor="w" if coluna == 0 else "center",
            ).grid(row=0, column=coluna, padx=14, pady=12, sticky="ew")

        self.linhas_relatorios = ctk.CTkScrollableFrame(
            self.tabela_relatorios, fg_color="transparent", corner_radius=0
        )
        self.linhas_relatorios.grid(row=1, column=0, columnspan=3, sticky="nsew")
        self.atualizar_relatorio_area()

    def obter_filtros_relatorio(self):

        ano_texto = self.combo_ano_relatorios.get()

        mes_texto = self.combo_mes_relatorios.get()

        semana_texto = self.combo_semana_relatorios.get()

        tipo = self.combo_tipo_relatorios.get()

        ano = int(ano_texto) if ano_texto != "Todos" else None

        mes = self.meses_relatorios.get(mes_texto)

        semanas = {
            "1ª": 1,
            "2ª": 2,
            "3ª": 3,
            "4ª": 4,
        }

        semana = semanas.get(semana_texto)

        return ano, mes, semana, tipo

    def alterar_filtro_relatorio(self, valor=None):

        self.atualizar_relatorio_area()

    def aplicar_filtros_relatorio(
        self,
        consulta,
        ano=None,
        mes=None,
        semana=None,
        tipo="Todas",
    ):

        data = Indicacao.data_indicacao

        if ano is not None:

            consulta = consulta.filter(extract("year", data) == ano)

        if mes is not None:

            consulta = consulta.filter(extract("month", data) == mes)

        if semana is not None:

            limites = {
                1: (1, 7),
                2: (8, 14),
                3: (15, 21),
                4: (22, 31),
            }

            inicio, fim = limites[semana]

            dia = extract("day", data)

            consulta = consulta.filter(
                dia >= inicio,
                dia <= fim,
            )

        if tipo != "Todas":

            consulta = consulta.filter(Indicacao.tipo_indicacao == tipo)

        return consulta

    def consultar_relatorio_area(self):

        ano, mes, semana, tipo = self.obter_filtros_relatorio()

        session = SessionLocal()

        try:

            areas = session.query(Area).order_by(Area.nome).all()

            resultados = []

            limites = {
                1: (1, 7),
                2: (8, 14),
                3: (15, 21),
                4: (22, 31),
            }

            for area in areas:

                nome_area = (area.nome or "").strip().casefold()

                if "projet" in nome_area:

                    consulta = session.query(
                        func.coalesce(func.max(Processo.quantidade), 0)
                    ).filter(Processo.area_id == area.id)

                    if ano is not None:

                        consulta = consulta.filter(
                            extract("year", Processo.data_processo) == ano
                        )

                    if mes is not None:

                        consulta = consulta.filter(
                            extract("month", Processo.data_processo) == mes
                        )

                    if semana is not None:

                        inicio, fim = limites[semana]

                        dia = extract("day", Processo.data_processo)

                        consulta = consulta.filter(
                            dia >= inicio,
                            dia <= fim,
                        )

                    total_processos = (consulta.scalar() or 0) if tipo == "Todas" else 0

                    total_indicacoes = 0

                else:

                    consulta = session.query(Indicacao).filter(
                        Indicacao.area_id == area.id
                    )

                    consulta = self.aplicar_filtros_relatorio(
                        consulta,
                        ano=ano,
                        mes=mes,
                        semana=semana,
                        tipo=tipo,
                    )

                    total_indicacoes = consulta.count()

                    total_processos = (
                        consulta.with_entities(
                            func.coalesce(func.max(Indicacao.quantidade), 0)
                        ).scalar()
                        or 0
                    )

                resultados.append(
                    (
                        area.nome,
                        int(total_processos),
                        total_indicacoes,
                    )
                )

            total_geral_processos = sum(processos for _, processos, _ in resultados)

            return resultados, total_geral_processos

        finally:

            session.close()

    def atualizar_relatorio_area(self):

        if not hasattr(self, "linhas_relatorios"):

            return

        for widget in self.linhas_relatorios.winfo_children():

            widget.destroy()

        try:

            resultados, total_processos = self.consultar_relatorio_area()

            self.card_relatorio_processos.atualizar_valor(total_processos)

            total_areas_com_resultados = sum(
                1
                for _, processos, indicacoes in resultados
                if processos > 0 or indicacoes > 0
            )

            self.card_relatorio_areas.atualizar_valor(total_areas_com_resultados)

            if not resultados:

                Mensagem(
                    self.linhas_relatorios,
                    "Nenhum resultado encontrado para os filtros selecionados.",
                    tipo="alerta",
                ).pack(pady=30)

                return

            for indice, (area, processos, indicacoes) in enumerate(resultados):

                linha = ctk.CTkFrame(
                    self.linhas_relatorios,
                    fg_color=(
                        CORES["fundo_secundario"] if indice % 2 == 0 else CORES["card"]
                    ),
                    corner_radius=6,
                )

                linha.pack(fill="x", pady=3)

                for coluna in range(3):

                    linha.grid_columnconfigure(
                        coluna,
                        weight=3 if coluna == 0 else 1,
                    )

                valores = [
                    area,
                    str(processos),
                    str(indicacoes),
                ]

                for coluna, valor in enumerate(valores):

                    ctk.CTkLabel(
                        linha,
                        text=valor,
                        text_color=CORES["texto"],
                        font=ctk.CTkFont(
                            size=13,
                            weight=("bold" if coluna == 0 else "normal"),
                        ),
                        anchor="w" if coluna == 0 else "center",
                    ).grid(
                        row=0,
                        column=coluna,
                        padx=12,
                        pady=10,
                        sticky="ew",
                    )

        except Exception as erro:

            Mensagem(
                self.linhas_relatorios,
                f"Erro ao gerar o relatório:\n{erro}",
                tipo="erro",
            ).pack(pady=30)

    def exportar_relatorio_excel(self):

        try:

            resultados, total_processos = self.consultar_relatorio_area()

            if not resultados:

                messagebox.showwarning(
                    "Relatório vazio",
                    "Não existem dados para exportar com os filtros selecionados.",
                )

                return

            ano, mes, semana, tipo = self.obter_filtros_relatorio()

            caminho = filedialog.asksaveasfilename(
                title="Salvar relatório por área",
                defaultextension=".xlsx",
                filetypes=[("Planilha Excel", "*.xlsx")],
                initialfile="relatorio_por_area.xlsx",
            )

            if not caminho:

                return

            workbook = Workbook()

            planilha = workbook.active

            planilha.title = "Relatório por Área"

            planilha.append(
                [
                    "Área",
                    "Total de Processos",
                    "Total de Indicações",
                ]
            )

            for area, processos, indicacoes in resultados:

                planilha.append(
                    [
                        area,
                        processos,
                        indicacoes,
                    ]
                )

            for celula in planilha[1]:

                celula.font = Font(
                    bold=True,
                    color="FFFFFF",
                )

                celula.fill = PatternFill(
                    fill_type="solid",
                    fgColor="244062",
                )

                celula.alignment = Alignment(horizontal="center")

            planilha.column_dimensions["A"].width = 30

            planilha.column_dimensions["B"].width = 24

            planilha.column_dimensions["C"].width = 24

            planilha.append([])

            planilha.append(["Filtros aplicados"])

            planilha.append(
                [
                    "Ano",
                    ano if ano is not None else "Todos",
                ]
            )

            planilha.append(
                [
                    "Mês",
                    self.combo_mes_relatorios.get(),
                ]
            )

            planilha.append(
                [
                    "Semana",
                    self.combo_semana_relatorios.get(),
                ]
            )

            planilha.append(
                [
                    "Tipo de indicação",
                    tipo,
                ]
            )

            planilha.append(
                [
                    "Total geral de processos",
                    total_processos,
                ]
            )

            workbook.save(caminho)

            messagebox.showinfo(
                "Exportação concluída",
                f"Relatório salvo com sucesso em:\n\n{caminho}",
            )

        except Exception as erro:

            messagebox.showerror(
                "Erro na exportação",
                f"Não foi possível exportar o relatório:\n\n{erro}",
            )
